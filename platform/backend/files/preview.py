import json
import mimetypes
from pathlib import Path

import pandas as pd
from fastapi.responses import JSONResponse, Response


PREVIEWABLE_TEXT = {
    ".txt", ".md", ".py", ".r", ".R", ".sql", ".log", ".sh",
    ".yaml", ".yml", ".toml", ".cfg", ".ini", ".env", ".example",
}
PREVIEWABLE_IMAGE = {".png", ".jpg", ".jpeg", ".gif", ".svg", ".webp"}
MAX_CSV_ROWS = 500
MAX_TEXT_BYTES = 1_000_000  # 1MB


def preview_file(file_path: Path) -> Response:
    ext = file_path.suffix.lower()

    if ext in PREVIEWABLE_IMAGE:
        return _preview_image(file_path, ext)
    elif ext == ".csv":
        return _preview_csv(file_path)
    elif ext == ".tsv":
        return _preview_csv(file_path, sep="\t")
    elif ext == ".json":
        return _preview_json(file_path)
    elif ext == ".html":
        return _preview_html(file_path)
    elif ext == ".docx":
        return _preview_docx(file_path)
    elif ext in PREVIEWABLE_TEXT:
        return _preview_text(file_path)
    else:
        return JSONResponse({"type": "unsupported", "name": file_path.name, "size": file_path.stat().st_size})


def _preview_image(path: Path, ext: str) -> Response:
    mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    data = path.read_bytes()
    return Response(content=data, media_type=mime)


def _preview_csv(path: Path, sep: str = ",") -> JSONResponse:
    try:
        df = pd.read_csv(path, sep=sep, nrows=MAX_CSV_ROWS, encoding="utf-8")
    except UnicodeDecodeError:
        df = pd.read_csv(path, sep=sep, nrows=MAX_CSV_ROWS, encoding="cp949")
    df = df.fillna("")
    return JSONResponse({
        "type": "csv",
        "columns": df.columns.tolist(),
        "rows": df.values.tolist(),
        "total_hint": "500+" if len(df) == MAX_CSV_ROWS else str(len(df)),
    })


def _preview_json(path: Path) -> JSONResponse:
    text = path.read_text(encoding="utf-8")[:MAX_TEXT_BYTES]
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        data = text
    return JSONResponse({"type": "json", "content": data})


def _preview_html(path: Path) -> JSONResponse:
    html = path.read_text(encoding="utf-8")
    return JSONResponse({"type": "html", "content": html})


def _preview_docx(path: Path) -> JSONResponse:
    try:
        from docx import Document
        from lxml import etree

        doc = Document(str(path))
        body = doc.element.body
        nsmap = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

        # Build lookup maps: xml element -> python-docx object (preserves order)
        para_map = {p._element: p for p in doc.paragraphs}
        table_map = {t._element: t for t in doc.tables}

        elements = []
        for child in body:
            tag = etree.QName(child.tag).localname
            if tag == "p" and child in para_map:
                para = para_map[child]
                text = para.text.strip() if para.text else ""
                if not text:
                    continue
                style = para.style.name if para.style else ""
                if style.startswith("Heading"):
                    try:
                        level = int(style.split()[-1])
                    except ValueError:
                        level = 1
                    elements.append({"kind": "heading", "level": level, "text": text})
                else:
                    elements.append({"kind": "paragraph", "text": text})
            elif tag == "tbl" and child in table_map:
                table = table_map[child]
                rows = []
                for row in table.rows:
                    rows.append([cell.text for cell in row.cells])
                elements.append({"kind": "table", "rows": rows})

        return JSONResponse({"type": "docx", "elements": elements})
    except Exception as e:
        return JSONResponse({"type": "error", "message": str(e)})


def _preview_text(path: Path) -> JSONResponse:
    try:
        text = path.read_text(encoding="utf-8")[:MAX_TEXT_BYTES]
    except UnicodeDecodeError:
        text = path.read_text(encoding="cp949")[:MAX_TEXT_BYTES]
    return JSONResponse({"type": "text", "content": text, "extension": path.suffix})
